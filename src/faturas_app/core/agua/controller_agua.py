"""Orquestração da extração de ÁGUA: lista PDFs, processa (cache JSON por PDF, progresso, cancelamento, paralelismo),
deduplica, aplica mapa/validação, monta os DataFrames das abas e grava/lê/concatena a planilha.

Uso (interface e CLI):
    lote = processar_pastas([pasta1, pasta2], progresso=cb, cache_dir=...)
    dfs = dataframes(lote)
    escrever_planilha(dfs, "faturas_agua.xlsx")
"""
from __future__ import annotations

import hashlib
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field

from ... import __version__
from .. import links
from . import derivados_agua, extrator, glossario_agua, mapa_conta
from . import schema_agua as S
from .base import Resultado

VERSAO_EXTRATOR = "4.0.0"      # entra na chave do cache: mudou o extrator → reprocessa
_LISTAS = ("faturas", "itens", "historico", "borderos", "contas_bordero", "avisos")


@dataclass
class Erro:
    arquivo: str
    mensagem: str


@dataclass
class ResultadoLote:
    resultado: Resultado = field(default_factory=Resultado)
    validacao: list = field(default_factory=list)
    erros: list = field(default_factory=list)
    duplicados: list = field(default_factory=list)      # (id, arquivo descartado)
    total: int = 0
    processados: int = 0
    do_cache: int = 0
    tempo: float = 0.0
    cancelado: bool = False

    def resumo(self) -> str:
        r = self.resultado
        fat = sum(1 for f in r.faturas if f.get("origem") != "analitica")
        ana = len(r.faturas) - fat
        erros = sum(1 for v in self.validacao if v["gravidade"] == "erro")
        avisos = sum(1 for v in self.validacao if v["gravidade"] == "aviso")
        ok = sum(1 for b in r.borderos if b.get("bate_total") == "SIM")
        txt = (f"{self.processados}/{self.total} PDF(s) em {self.tempo:,.0f}s"
               + (f" ({self.do_cache} do cache)" if self.do_cache else "")
               + f" · {fat} fatura(s) · {ana} conta(s) de analítica · {len(r.borderos)} borderô(s)"
               + (f" ({ok} conferem)" if r.borderos else "")
               + f" · validação: {erros} erro(s), {avisos} aviso(s)"
               + (f" · {len(self.erros)} PDF(s) com falha" if self.erros else "")
               + (" · CANCELADO" if self.cancelado else ""))
        return txt


# ── listagem ──────────────────────────────────────────────────────────────────
def listar_pdfs(pasta: str, subpastas: bool = True) -> list[str]:
    saida = []
    if not pasta or not os.path.isdir(pasta):
        return saida
    if subpastas:
        for raiz, _d, arqs in os.walk(pasta):
            saida += [os.path.join(raiz, a) for a in arqs if a.lower().endswith(".pdf")]
    else:
        saida = [os.path.join(pasta, a) for a in os.listdir(pasta) if a.lower().endswith(".pdf")]
    return sorted(saida)


def contar_pdfs(pastas: list[str], subpastas: bool = True) -> int:
    return sum(len(listar_pdfs(p, subpastas)) for p in pastas)


# ── cache ─────────────────────────────────────────────────────────────────────
def _chave_cache(pdf: str) -> str:
    st = os.stat(pdf)
    bruto = f"{VERSAO_EXTRATOR}|{os.path.abspath(pdf).lower()}|{st.st_size}|{st.st_mtime_ns}"
    return hashlib.sha1(bruto.encode("utf-8")).hexdigest()


def _res_para_dict(res: Resultado) -> dict:
    return {k: getattr(res, k) for k in _LISTAS}


def _dict_para_res(d: dict) -> Resultado:
    res = Resultado()
    for k in _LISTAS:
        getattr(res, k).extend(d.get(k) or [])
    return res


def _processar_um(pdf: str, cache_dir: str | None, ocr: bool):
    """Devolve (pdf, Resultado|None, erro|None, veio_do_cache)."""
    try:
        fp = os.path.join(cache_dir, _chave_cache(pdf) + ".json") if cache_dir else None
        if fp and os.path.isfile(fp):
            with open(fp, "r", encoding="utf-8") as f:
                return pdf, _dict_para_res(json.load(f)), None, True
        res = extrator.processar_pdf(pdf, ocr=ocr)
        if fp:
            tmp = fp + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(_res_para_dict(res), f, ensure_ascii=False)
            os.replace(tmp, fp)
        return pdf, res, None, False
    except Exception as e:  # noqa: BLE001 — registrado no lote, o resto segue
        return pdf, None, f"{type(e).__name__}: {e}", False


# ── lote ──────────────────────────────────────────────────────────────────────
def processar_pastas(pastas: list[str], progresso=None, cancelar=None, cache_dir: str | None = None, ocr: bool = True,
                     paralelo: int | None = None, subpastas: bool = True, modo_link: str = links.MODO_BUSCA,
                     template: str | None = None, arquivos: list[str] | None = None) -> ResultadoLote:
    """Processa todos os PDFs das pastas. `progresso(i, n, nome)` a cada PDF; `cancelar()` → True interrompe."""
    pdfs = list(arquivos) if arquivos is not None else []
    if arquivos is None:
        for p in pastas:
            pdfs += listar_pdfs(p, subpastas)
    lote = ResultadoLote(total=len(pdfs))
    if cache_dir:
        os.makedirs(cache_dir, exist_ok=True)
    t0 = time.time()
    n_workers = max(1, paralelo or min(4, (os.cpu_count() or 2) - 1))
    parciais: dict[str, tuple] = {}

    def _andamento(nome):
        lote.processados += 1
        if progresso:
            progresso(lote.processados, lote.total, nome)

    if n_workers == 1 or len(pdfs) <= 1:
        for pdf in pdfs:
            if cancelar and cancelar():
                lote.cancelado = True
                break
            parciais[pdf] = _processar_um(pdf, cache_dir, ocr)
            _andamento(os.path.basename(pdf))
    else:
        with ThreadPoolExecutor(max_workers=n_workers) as ex:
            futuros = {ex.submit(_processar_um, pdf, cache_dir, ocr): pdf for pdf in pdfs}
            for fut in as_completed(futuros):
                pdf = futuros[fut]
                parciais[pdf] = fut.result()
                _andamento(os.path.basename(pdf))
                if cancelar and cancelar():
                    lote.cancelado = True
                    for f2 in futuros:
                        f2.cancel()
                    break

    # junta na ordem dos PDFs, deduplicando por id (faturas/borderôs) e levando itens/histórico/contas junto
    vistos_f: set[str] = set()
    vistos_b: set[str] = set()
    for pdf in pdfs:
        item = parciais.get(pdf)
        if not item:
            continue
        _pdf, res, erro, do_cache = item
        nome = os.path.basename(pdf)
        lote.do_cache += bool(do_cache)
        if erro:
            lote.erros.append(Erro(arquivo=nome, mensagem=erro))
            continue
        link = links.gerar_link(modo_link, nome, pdf, template)
        ids_ok: set[str] = set()
        for f in res.faturas:
            if f["id_fatura_agua"] in vistos_f:
                lote.duplicados.append((f["id_fatura_agua"], nome))
                continue
            vistos_f.add(f["id_fatura_agua"])
            ids_ok.add(f["id_fatura_agua"])
            f["link_pdf"] = link
            lote.resultado.faturas.append(f)
        lote.resultado.itens.extend(i for i in res.itens if i.get("id_fatura_agua") in ids_ok)
        lote.resultado.historico.extend(h for h in res.historico if h.get("id_fatura_agua") in ids_ok)
        ids_b: set[str] = set()
        for b in res.borderos:
            if b["id_bordero_agua"] in vistos_b:
                lote.duplicados.append((b["id_bordero_agua"], nome))
                continue
            vistos_b.add(b["id_bordero_agua"])
            ids_b.add(b["id_bordero_agua"])
            lote.resultado.borderos.append(b)
        lote.resultado.contas_bordero.extend(c for c in res.contas_bordero if c.get("id_bordero_agua") in ids_b)
        lote.resultado.avisos.extend(res.avisos)
    lote.validacao = derivados_agua.aplicar(lote.resultado, lote.duplicados)
    lote.tempo = time.time() - t0
    return lote


# ── DataFrames / planilha ─────────────────────────────────────────────────────
def dataframes(lote: ResultadoLote | Resultado, validacao: list | None = None, incluir_mapa: bool = True) -> dict:
    import pandas as pd

    res = lote.resultado if isinstance(lote, ResultadoLote) else lote
    val = lote.validacao if isinstance(lote, ResultadoLote) else (validacao or [])
    dados = {
        S.ABA_FATURAS: res.faturas, S.ABA_ITENS: res.itens, S.ABA_HISTORICO: res.historico,
        S.ABA_BORDEROS: res.borderos, S.ABA_CONTAS_BORDERO: res.contas_bordero, S.ABA_VALIDACAO: val,
    }
    dfs = {}
    for aba, linhas in dados.items():
        cols = S.CANONICAL_COLUMNS[aba]
        df = pd.DataFrame([{c: r.get(c) for c in cols} for r in linhas], columns=cols)
        for c in S.COLUNAS_NUMERICAS.get(aba, []):
            df[c] = pd.to_numeric(df[c], errors="coerce")
        if aba == S.ABA_FATURAS:
            df = df.sort_values(["fornecedor", "competencia", "conta"], na_position="last", kind="stable").reset_index(drop=True)
        dfs[aba] = df
    if incluir_mapa and mapa_conta.ativo():
        regs = mapa_conta.registros()
        dfs[S.ABA_MAPA] = pd.DataFrame([{c: (", ".join(r.get(c)) if c == "identificadores" and isinstance(r.get(c), list) else r.get(c))
                                         for c in S.MAPA_COLS} for r in regs], columns=S.MAPA_COLS)
    return dfs


def _estilizar(caminho: str, abas: list[str]) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(border_style="thin", color="BFBFBF")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = load_workbook(caminho)
    for sn in abas:
        if sn not in wb.sheetnames:
            continue
        hc, rc = S.SHEET_COLORS.get(sn, ("404040", "D9D9D9"))
        ws = wb[sn]
        hf, rf = PatternFill("solid", fgColor=hc), PatternFill("solid", fgColor=rc)
        for cell in ws[1]:
            cell.fill = hf
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = brd
        ws.row_dimensions[1].height = 30
        alt = False
        for row in ws.iter_rows(min_row=2):
            f = PatternFill("solid", fgColor="FFFFFF") if not alt else rf
            for c in row:
                c.fill, c.font, c.border = f, Font(size=10), brd
                c.alignment = Alignment(vertical="center")
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
            alt = not alt
        for col in ws.columns:
            mx = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 4, 55)
        ws.freeze_panes = "A2"
    wb.save(caminho)


def escrever_planilha(dfs: dict, caminho: str) -> str:
    """Grava as abas (na ordem canônica + glossário) com a formatação padrão do app."""
    import pandas as pd

    dfs = glossario_agua.garantir_glossario(dfs)
    abas = [a for a in S.SHEET_ORDER if a in dfs] + [a for a in dfs if a not in S.SHEET_ORDER]
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        for aba in abas:
            df = dfs[aba]
            if df is None or df.empty:
                pd.DataFrame({"(sem dados)": []}).to_excel(w, sheet_name=aba, index=False)
            else:
                df.to_excel(w, sheet_name=aba, index=False)
    _estilizar(caminho, abas)
    return caminho


def gravar_csv(dfs: dict, pasta: str) -> None:
    os.makedirs(pasta, exist_ok=True)
    for aba, df in dfs.items():
        if df is None or df.empty:
            continue
        df.to_csv(os.path.join(pasta, f"{aba}.csv"), sep=";", decimal=",", index=False, encoding="utf-8-sig")


def ler_planilha(caminho: str) -> dict:
    import pandas as pd

    xls = pd.ExcelFile(caminho, engine="openpyxl")
    dfs = {}
    for nome in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=nome, dtype=object)
        if list(df.columns) == ["(sem dados)"]:
            df = pd.DataFrame()
        dfs[nome] = df
    return dfs


def divergencias(base_dfs: dict) -> list[str]:
    """A planilha base precisa ter as abas de água com as colunas canônicas."""
    msgs = []
    if S.ABA_FATURAS not in base_dfs and S.ABA_BORDEROS not in base_dfs:
        return [f"A planilha não tem as abas de água ('{S.ABA_FATURAS}' / '{S.ABA_BORDEROS}')."]
    for aba in (S.ABA_FATURAS, S.ABA_ITENS, S.ABA_HISTORICO, S.ABA_BORDEROS, S.ABA_CONTAS_BORDERO):
        df = base_dfs.get(aba)
        if df is None or df.empty:
            continue
        faltantes = [c for c in S.CANONICAL_COLUMNS[aba] if c not in df.columns]
        if faltantes:
            msgs.append(f"Aba '{aba}': coluna(s) ausente(s) — {', '.join(faltantes)}.")
    return msgs


def concatenar(base_dfs: dict, novos_dfs: dict) -> tuple[dict, list[str]]:
    """Base + novos, deduplicando por id (faturas/borderôs) e por linha inteira nas demais abas."""
    import pandas as pd

    resultado, resumo = {}, []
    for aba in S.SHEET_ORDER:
        cols = S.CANONICAL_COLUMNS[aba]
        base = base_dfs.get(aba)
        base = base.reindex(columns=cols) if base is not None and not base.empty else pd.DataFrame(columns=cols)
        novo = novos_dfs.get(aba)
        novo = novo.reindex(columns=cols) if novo is not None and not novo.empty else pd.DataFrame(columns=cols)
        if aba == S.ABA_MAPA:
            comb = novo if not novo.empty else base
            resultado[aba] = comb.reset_index(drop=True)
            continue
        comb = pd.concat([base, novo], ignore_index=True)
        antes = len(comb)
        chave = S.DEDUP_KEYS.get(aba)
        comb = comb.drop_duplicates(subset=chave, keep="first").reset_index(drop=True)
        for c in S.COLUNAS_NUMERICAS.get(aba, []):
            comb[c] = pd.to_numeric(comb[c], errors="coerce")
        resultado[aba] = comb
        linha = f"Aba '{aba}': {len(base)} (base) + {len(novo)} (novos) = {len(comb)} linha(s)."
        if antes - len(comb):
            linha += f" ({antes - len(comb)} duplicata(s) removida(s))."
        resumo.append(linha)
    for aba, df in base_dfs.items():
        if aba not in resultado and df is not None and not df.empty and not str(aba).lower().startswith("gloss"):
            resultado[aba] = df
    return resultado, resumo


def salvar_mapa_da_extracao(res: Resultado) -> dict:
    """Constrói o mapa de contas a partir do resultado e o salva como mapa ativo."""
    regs = mapa_conta.construir_de_extracao(res.faturas, res.contas_bordero)
    return mapa_conta.salvar(regs, origem="extração das faturas")
